"""Generate the binary test documents (PDF, DOCX) in testdata/reports/.

All content is fictional and non-confidential (R3).
Requires the dev extras: pip install -e ".[dev]"
"""

from pathlib import Path

HERE = Path(__file__).parent
REPORTS = HERE / "reports"

PDF_PAGES = [
    (
        "Kelp-3 Buoy Field Trial Report 2025",
        [
            "Brightwater Labs - Hardware Engineering. Fictional document for testing.",
            "Executive Summary",
            "The Kelp-3 buoy was trialled for 120 days between April and July 2025 at four sites "
            "in the North Sea and the Baltic. The trial met its main goal: average uptime was 99.1 percent, "
            "above the 98 percent target. The main weakness was the anchor chain swivel, which failed on two buoys.",
            "We recommend a production run of 200 units in 2026, conditional on a redesigned swivel.",
        ],
    ),
    (
        "1. Trial Setup",
        [
            "Twelve Kelp-3 buoys were deployed: four at Helgoland, three at Fehmarn, three at Rugen and two at "
            "Bornholm. Each buoy carried a wave sensor, a water-temperature string with 8 sensors down to 20 m, "
            "and a solar charging system with a 240 Wh lithium iron phosphate battery.",
            "Data was transmitted every 10 minutes over a satellite link. When the link was unavailable, readings "
            "were buffered on board for up to 14 days.",
        ],
    ),
    (
        "2. Power and Battery Results",
        [
            "The solar panels produced on average 410 Wh per day in May and 520 Wh per day in June. The lowest "
            "daily yield, 95 Wh, was recorded at Bornholm during a storm on 14 April.",
            "Battery state of charge never dropped below 38 percent. Battery capacity loss after 120 days was "
            "1.6 percent, which projects to a battery service life of about seven years.",
        ],
    ),
    (
        "3. Data Quality",
        [
            "Wave height readings were compared against a reference Waverider buoy at Helgoland. The mean "
            "absolute difference was 0.07 m. Water-temperature sensors drifted by less than 0.05 degrees Celsius.",
            "Biofouling on the temperature string became visible after 60 days. An anti-fouling copper sleeve "
            "is recommended for deployments longer than three months.",
        ],
    ),
    (
        "4. Failures and Incidents",
        [
            "Two anchor chain swivels failed due to corrosion, at Fehmarn on day 71 and at Rugen on day 93. "
            "Both buoys drifted but were recovered within 36 hours thanks to the GPS tracker.",
            "One satellite modem froze and needed a remote power cycle. Firmware version 2.4.1 fixed the cause, "
            "a watchdog timer that was disabled by mistake.",
        ],
    ),
    (
        "5. Costs and Recommendations",
        [
            "The bill of materials per buoy was 4,850 euros, 6 percent above the estimate, mainly because of "
            "the satellite modem price. Deployment cost per buoy including vessel time was 1,900 euros.",
            "Recommendations: (1) replace the galvanised swivel with a duplex stainless steel swivel; "
            "(2) add the copper anti-fouling sleeve; (3) ship firmware 2.4.1 or newer; "
            "(4) proceed with a production run of 200 units in 2026.",
        ],
    ),
]


def make_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_title("Kelp-3 Buoy Field Trial Report 2025")
    pdf.set_author("Brightwater Labs (fictional)")
    for heading, paragraphs in PDF_PAGES:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.multi_cell(0, 10, heading)
        pdf.ln(4)
        for para in paragraphs:
            bold = len(para) < 40
            pdf.set_font("Helvetica", "B" if bold else "", 13 if bold else 11)
            pdf.multi_cell(0, 6, para)
            pdf.ln(3)
    pdf.output(str(path))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = "Architecture Review: Forecast Pipeline v3"
    d.add_heading("Architecture Review: Forecast Pipeline v3", level=0)
    d.add_paragraph("Fictional document for testing. Review date: 2025-09-12. Reviewers: platform guild.")
    d.add_heading("Context", level=1)
    d.add_paragraph(
        "The current forecast pipeline (v2) recomputes all station forecasts every three hours in a single "
        "batch job. With 1,400 stations the batch takes 95 minutes, leaving little headroom."
    )
    d.add_heading("Proposal", level=1)
    d.add_paragraph(
        "Pipeline v3 splits the batch into per-region jobs running on a queue. Each job reads buoy readings "
        "from the time-series store, runs the wave model and writes forecasts to the forecast cache."
    )
    d.add_heading("Queue Technology", level=2)
    d.add_paragraph(
        "The team evaluated three queues and chose the managed message queue 'Current' because it supports "
        "message deduplication and delayed retries. A self-hosted option was rejected due to operating cost."
    )
    d.add_heading("Decisions", level=1)
    d.add_paragraph("Decision 1: Forecast jobs must be idempotent; re-running a job for the same region and hour "
                    "overwrites the previous result.")
    d.add_paragraph("Decision 2: The target runtime for a full cycle is 30 minutes for 3,000 stations.")
    d.add_paragraph("Decision 3: The v2 batch job stays available as a fallback until March 2026.")
    d.add_heading("Risks", level=1)
    table = d.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Risk", "Impact", "Mitigation"
    for row in [
        ("Queue backlog during storms", "High", "Autoscale workers on queue depth above 500 messages"),
        ("Inconsistent forecasts across regions", "Medium", "Version every forecast with the model run ID"),
    ]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.save(str(path))


if __name__ == "__main__":
    REPORTS.mkdir(parents=True, exist_ok=True)
    make_pdf(REPORTS / "kelp3-field-trial-2025.pdf")
    make_docx(REPORTS / "forecast-pipeline-v3-review.docx")
    print("wrote", *sorted(p.name for p in REPORTS.iterdir()))
